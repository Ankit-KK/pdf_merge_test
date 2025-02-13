import streamlit as st
import fitz  # PyMuPDF
from pptx import Presentation
from pptx.util import Inches
from io import BytesIO

def merge_pdf_pages(pdf_files):
    merged_pdf = fitz.open()
    for pdf_file in pdf_files:
        pdf_document = fitz.open(stream=pdf_file.read(), filetype="pdf")
        merged_pdf.insert_pdf(pdf_document, from_page=0, to_page=3)
    output = BytesIO()
    merged_pdf.save(output)
    return output.getvalue()

def merge_ppt_slides(ppt_files):
    merged_presentation = Presentation()
    for ppt_file in ppt_files:
        presentation = Presentation(BytesIO(ppt_file.read()))
        for slide in presentation.slides:
            merged_presentation.slides.add_slide(slide.slide_layout)
    output = BytesIO()
    merged_presentation.save(output)
    return output.getvalue()

st.title("PDF and PPT Page Merger")

uploaded_files = st.file_uploader("Upload PDF or PPT files", type=["pdf", "pptx"], accept_multiple_files=True)

if uploaded_files:
    pdf_files = [file for file in uploaded_files if file.name.endswith('.pdf')]
    ppt_files = [file for file in uploaded_files if file.name.endswith('.pptx')]

    if pdf_files:
        st.write("Merging PDF pages...")
        merged_pdf = merge_pdf_pages(pdf_files)
        st.download_button("Download Merged PDF", merged_pdf, "merged.pdf")

    if ppt_files:
        st.write("Merging PPT slides...")
        merged_ppt = merge_ppt_slides(ppt_files)
        st.download_button("Download Merged PPT", merged_ppt, "merged.pptx")
